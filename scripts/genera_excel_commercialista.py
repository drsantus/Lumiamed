"""
Crea docs/LumiaMed_verifica_commercialista.xlsx a partire da docs/casi-test.json.

Uso:
  npm run casi-test
  python scripts/genera_excel_commercialista.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATI = json.loads((ROOT / "docs" / "casi-test.json").read_text(encoding="utf-8"))
OUT = ROOT / "docs" / "LumiaMed_verifica_commercialista.xlsx"

FONT = "Arial"
H_FILL = PatternFill("solid", start_color="1F3A5F")
H_FONT = Font(name=FONT, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
BODY = Font(name=FONT, size=10)
BLUE = Font(name=FONT, size=10, color="0000FF")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
EURO = '#,##0.00 "€";-#,##0.00 "€";"-"'
PCT = "0.00%"

wb = Workbook()

# ---------------- Foglio 1: istruzioni ----------------
ist = wb.active
ist.title = "Istruzioni"
ist.column_dimensions["A"].width = 110
testo = [
    ("LumiaMed — verifica dei calcoli fiscali", Font(name=FONT, bold=True, size=14)),
    ("", BODY),
    ("Gentile dottore/dottoressa commercialista,", BODY),
    ("questo file contiene i casi di prova del motore di calcolo dell'app LumiaMed (gestione finanziaria per medici liberi professionisti).", BODY),
    ("Le chiediamo di verificare se i valori calcolati dall'app sono corretti.", BODY),
    ("", BODY),
    ("Come compilare (solo le celle GIALLE):", Font(name=FONT, bold=True)),
    ("1. Foglio 'Casi di prova': per ogni riga scelga 'Sì' o 'No' nella colonna 'Corretto?'.", BODY),
    ("2. Se è 'No', scriva il valore giusto in 'Valore corretto' e il motivo in 'Note'. La colonna 'Differenza' si calcola da sola.", BODY),
    ("3. Foglio 'Parametri': confermi o corregga i parametri normativi 2026 usati dall'app.", BODY),
    ("", BODY),
    ("Note importanti:", Font(name=FONT, bold=True)),
    ("• Tutti i calcoli seguono il principio di cassa (ricavi incassati e spese pagate nell'anno).", BODY),
    ("• Le addizionali dei casi 'Ordinario' usano tabelle DI PROVA (non quelle reali di un comune): serve verificare il metodo di calcolo.", BODY),
    ("• Quota A ENPAM: l'app applica la fascia d'età superiore dal mese successivo al compleanno (pro-rata mensile).", BODY),
    ("• Quota B ENPAM: base = reddito libero-professionale (forfettario: ricavi × 78%, rigo LM34; ordinario: compensi − spese), meno la soglia già coperta dalla Quota A.", BODY),
    ("• Acconti: metodo storico, soggetti ISA/forfettari in due rate da 50%.", BODY),
    ("• Bollo da 2 € riaddebitato al paziente: l'app ora lo ESCLUDE dal reddito (D.Lgs. 192/2024). Da confermare.", BODY),
    ("", BODY),
    ("Esempio di compilazione: riga con valore errato → Corretto? = No; Valore corretto = 1.450,00; Note = 'manca X'.", Font(name=FONT, italic=True, size=10)),
]
for i, (t, f) in enumerate(testo, start=1):
    c = ist.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True)

# ---------------- Foglio 2: casi di prova ----------------
ws = wb.create_sheet("Casi di prova")
headers = ["ID", "Area", "Scenario", "Dati di partenza", "Voce", "Valore calcolato dall'app", "Corretto?", "Valore corretto", "Differenza", "Note del commercialista"]
widths = [6, 15, 42, 60, 38, 18, 11, 16, 14, 40]
for col, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=1, column=col, value=h)
    c.font, c.fill, c.alignment, c.border = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center"), BORDER
    ws.column_dimensions[get_column_letter(col)].width = w
ws.row_dimensions[1].height = 32
ws.freeze_panes = "A2"

dv = DataValidation(type="list", formula1='"Sì,No"', allow_blank=True)
ws.add_data_validation(dv)

for r, riga in enumerate(DATI["righe"], start=2):
    valori = [riga["id"], riga["area"], riga["descrizione"], riga["input"], riga["voce"]]
    for col, v in enumerate(valori, start=1):
        c = ws.cell(row=r, column=col, value=v)
        c.font, c.alignment, c.border = BODY, WRAP, BORDER
    vc = ws.cell(row=r, column=6)
    if riga["formato"] == "testo":
        vc.value = riga.get("testo", "")
    else:
        vc.value = riga["valore"]
        vc.number_format = EURO if riga["formato"] == "euro" else PCT
    vc.font, vc.border, vc.alignment = BLUE, BORDER, Alignment(vertical="top")
    for col in (7, 8, 10):
        c = ws.cell(row=r, column=col)
        c.fill, c.border, c.font, c.alignment = INPUT_FILL, BORDER, BODY, WRAP
    ws.cell(row=r, column=8).number_format = EURO if riga["formato"] == "euro" else (PCT if riga["formato"] == "percentuale" else "General")
    dv.add(f"G{r}")
    d = ws.cell(row=r, column=9)
    if riga["formato"] == "testo":
        d.value = f'=IF(H{r}="","",IF(H{r}=F{r},"uguale","diverso"))'
    else:
        d.value = f'=IF(H{r}="","",H{r}-F{r})'
        d.number_format = EURO if riga["formato"] == "euro" else PCT
    d.font, d.border, d.alignment = BODY, BORDER, Alignment(vertical="top")

last = len(DATI["righe"]) + 1
ws.conditional_formatting.add(f"A2:J{last}", FormulaRule(formula=[f'$G2="No"'], fill=PatternFill("solid", start_color="F8D7DA")))
ws.conditional_formatting.add(f"A2:J{last}", FormulaRule(formula=[f'$G2="Sì"'], fill=PatternFill("solid", start_color="D4EDDA")))
ws.auto_filter.ref = f"A1:J{last}"

riep = last + 2
ws.cell(row=riep, column=5, value="Righe verificate (Sì + No)").font = Font(name=FONT, bold=True)
ws.cell(row=riep, column=6, value=f'=COUNTIF(G2:G{last},"Sì")+COUNTIF(G2:G{last},"No")').font = BODY
ws.cell(row=riep + 1, column=5, value="Righe da correggere (No)").font = Font(name=FONT, bold=True)
ws.cell(row=riep + 1, column=6, value=f'=COUNTIF(G2:G{last},"No")').font = BODY
ws.cell(row=riep + 2, column=5, value="Righe totali").font = Font(name=FONT, bold=True)
ws.cell(row=riep + 2, column=6, value=f"=COUNTA(A2:A{last})").font = BODY

# ---------------- Foglio 3: parametri ----------------
ps = wb.create_sheet("Parametri")
ph = ["Gruppo", "Parametro", "Valore usato dall'app", "Fonte", "Nota", "Confermato?", "Valore corretto / commento"]
pw = [22, 48, 34, 55, 50, 12, 40]
for col, (h, w) in enumerate(zip(ph, pw), start=1):
    c = ps.cell(row=1, column=col, value=h)
    c.font, c.fill, c.alignment, c.border = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center"), BORDER
    ps.column_dimensions[get_column_letter(col)].width = w
ps.freeze_panes = "A2"
dv2 = DataValidation(type="list", formula1='"Sì,No"', allow_blank=True)
ps.add_data_validation(dv2)
for r, par in enumerate(DATI["parametri"], start=2):
    vals = [par["gruppo"], par["voce"], par["valore"], par["fonte"], par.get("nota") or ""]
    for col, v in enumerate(vals, start=1):
        c = ps.cell(row=r, column=col, value=v)
        c.font = BLUE if col == 3 else BODY
        c.alignment, c.border = WRAP, BORDER
    for col in (6, 7):
        c = ps.cell(row=r, column=col)
        c.fill, c.border, c.font, c.alignment = INPUT_FILL, BORDER, BODY, WRAP
    dv2.add(f"F{r}")
plast = len(DATI["parametri"]) + 1
ps.conditional_formatting.add(f"A2:G{plast}", FormulaRule(formula=[f'$F2="No"'], fill=PatternFill("solid", start_color="F8D7DA")))
ps.auto_filter.ref = f"A1:G{plast}"
ps.cell(row=2, column=5).comment = Comment("Le note segnalano i punti dove le fonti pubbliche non concordano.", "LumiaMed")

wb.save(OUT)
print(f"Creato {OUT}")
